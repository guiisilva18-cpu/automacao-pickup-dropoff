"""Extrai a "Taxa de coleta no prazo" (JMS > Indicadores de Negócios > Prazo >
Taxa de coleta no prazo) para um conjunto fixo de bases, gera uma planilha
diária (uma linha por base) e manda por e-mail.

Mesmo endpoint que Automacao_JMS/extrair_taxa_coleta.py já usa
(timely_collection_rate_collect_network, conta do Vanilson via
JMS_TOKEN_INDICADORES), mas com 3 diferenças propositais:

1. timeType="2" (Horário de término do prazo de coleta) em vez de "1" —
   pedido explícito do Guilherme (2026-08-11): esse relatório usa o horário
   de término do prazo, não o de início.
2. Roda sempre pra D-1 (dia anterior) — o dia corrente ainda está em
   andamento, o número só fecha depois da virada.
3. Só interessa um conjunto fixo de 14 bases (não a rede toda) e captura 2
   campos que o script original ignora (timelyTakingNum, shouldTakingNum
   isolado por base) pra bater com as 4 colunas da planilha de referência
   (print do Guilherme, 2026-08-11): Qtd a coletar / Qtd coletada no prazo /
   Soma coletados+tentativas / Taxa de coleta Real.

Descoberta dos campos (captura de rede, 2026-08-11): cada registro já vem
com shouldTakingNum, timelyTakingNum, timelyTryTakingNum e
timelyPickRateTotal por base — não precisa de chamada nova nem de uma
chamada por base.
"""
import json
import logging
import os
import smtplib
import ssl
import sys
from collections import defaultdict
from datetime import date, timedelta
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()

PASTA_BASE = Path(__file__).parent
PASTA_LOGS = PASTA_BASE / "logs"
PASTA_OUTPUT = PASTA_BASE / "output"
PASTA_DADOS = PASTA_BASE / "dados"
PASTA_LOGS.mkdir(exist_ok=True)
PASTA_OUTPUT.mkdir(exist_ok=True)
PASTA_DADOS.mkdir(exist_ok=True)

# Histórico acumulado do mês (mesmo motivo do extrair_dropoff.py,
# 2026-08-11): cada dia roda 1x e salva aqui, o painel mensal é montado a
# partir desse arquivo em vez de rebuscar o mês inteiro na API toda vez.
HISTORICO_PATH = PASTA_DADOS / "historico_pickup.json"


def carregar_historico() -> dict[str, list[dict]]:
    if not HISTORICO_PATH.exists():
        return {}
    return json.loads(HISTORICO_PATH.read_text(encoding="utf-8"))


def salvar_historico(historico: dict[str, list[dict]]):
    HISTORICO_PATH.write_text(
        json.dumps(historico, ensure_ascii=False, indent=2), encoding="utf-8"
    )

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(PASTA_LOGS / "extrair_pickup.log", encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger(__name__)

URL_TAXA_COLETA = (
    "https://gw.jtjms-br.com/businessindicator/bigdataReport/detailDir/"
    "datacabin/composite_waybill/timely_collection_rate_collect_network"
)

# Relatório de Monitoramento EPOP — traz 1 linha por loja/pedido do dia, com
# `signaturePictureUrl` (caminho do arquivo quando tem foto de comprovação
# da coleta, null quando não tem). Usado só pra calcular "Transferência
# POC" = % de linhas da base que tem foto (pedido do Guilherme, 2026-08-11).
# URL real capturada por rede (o nome "bus_epop_monitor_count" que aparece
# no DevTools é só um apelido, a URL completa é essa aqui).
URL_EPOP = "https://gw.jtjms-br.com/businessindicator/bigdataReport/detail/bus_epop_monitor_count"

HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Content-Type": "application/json;charset=UTF-8",
    "lang": "PT",
    "langType": "PT",
    "routeName": "CollectTimelyRate",
}

TAMANHO_PAGINA = 1000
ORIGEM_PEDIDO_FILTRO = "TikTok"

# Bases do relatório PICKUP (Guilherme, 2026-08-11) — só essas contam, o
# resto do retorno da API (outras ~20 bases da rede) é descartado.
BASES_PICKUP = [
    "CARAP 02-SP",
    "CARAP-SP",
    "CHM-SP",
    "CLP-SP",
    "COT-SP",
    "F JND-SP",
    "F S-JRG-SP",
    "ITUP-SP",
    "JND 02-SP",
    "JND-SP",
    "OSC 02-SP",
    "OSC-SP",
    "S-CSVD-SP",
    "S-FREG-SP",
]


class TokenExpiradoError(Exception):
    pass


def _buscar_registros_brutos(dia: str) -> list[dict]:
    token = os.environ["JMS_TOKEN_INDICADORES"]
    headers = {**HEADERS, "authToken": token}

    registros = []
    current = 1
    while True:
        payload = {
            "current": current,
            "size": TAMANHO_PAGINA,
            "startTime": f"{dia} 00:00:00",
            "endTime": f"{dia} 23:59:59",
            "timeType": "2",
            "typeId": 0,
            "countryId": "1",
        }
        resp = requests.post(URL_TAXA_COLETA, headers=headers, json=payload, timeout=30)
        if resp.status_code in (401, 403):
            raise TokenExpiradoError(
                "JMS_TOKEN_INDICADORES expirado ou sem sessão ativa. "
                "Faça login no portal (conta do Vanilson) e atualize JMS_TOKEN_INDICADORES no .env."
            )
        resp.raise_for_status()
        resultado = resp.json()
        if resultado.get("code") != 1:
            raise RuntimeError(f"Erro ao buscar taxa de coleta: {resultado}")

        dados = resultado.get("data") or {}
        pagina = dados.get("records", [])
        registros.extend(pagina)
        total = dados.get("total", 0)
        if not pagina or len(registros) >= total:
            break
        current += 1
        if current > 5:
            log.warning("Parou de paginar em 5 páginas (total esperado: %s)", total)
            break

    return registros


def _buscar_registros_epop_brutos(dia: str) -> list[dict]:
    """Pagina o Relatório de Monitoramento EPOP pro dia inteiro. Tamanho
    máximo de página é 1000 (confirmado por erro da própria API em teste:
    "size过大，不能超过1000") — mesmo teto do endpoint de Taxa de Coleta."""
    token = os.environ["JMS_TOKEN_INDICADORES"]
    headers = {**HEADERS, "authToken": token}

    registros = []
    current = 1
    while True:
        payload = {
            "current": current,
            "size": TAMANHO_PAGINA,
            "groupByKey": "network",
            "startTime": f"{dia} 00:00:00",
            "endTime": f"{dia} 23:59:59",
            "countryId": "1",
        }
        resp = requests.post(URL_EPOP, headers=headers, json=payload, timeout=30)
        if resp.status_code in (401, 403):
            raise TokenExpiradoError(
                "JMS_TOKEN_INDICADORES expirado ou sem sessão ativa. "
                "Faça login no portal (conta do Vanilson) e atualize JMS_TOKEN_INDICADORES no .env."
            )
        resp.raise_for_status()
        resultado = resp.json()
        if resultado.get("code") != 1:
            raise RuntimeError(f"Erro ao buscar relatório EPOP: {resultado}")

        dados = resultado.get("data") or {}
        pagina = dados.get("records", [])
        registros.extend(pagina)
        total = dados.get("total", 0)
        if not pagina or len(registros) >= total:
            break
        current += 1
        if current > 20:
            log.warning("EPOP: parou de paginar em 20 páginas (total esperado: %s)", total)
            break

    return registros


def buscar_pickup(dia: str) -> list[dict]:
    """Agrega os registros brutos por base (só ORIGEM_PEDIDO_FILTRO) e
    filtra pra só as bases em BASES_PICKUP. Loga um aviso pra qualquer base
    da lista que não apareceu na resposta da API."""
    brutos = _buscar_registros_brutos(dia)

    agregado: dict[str, dict[str, int]] = defaultdict(
        lambda: {"deveria": 0, "coletado_no_prazo": 0, "tentativas": 0}
    )
    for r in brutos:
        base = (r.get("pickNetworkName") or "").strip()
        if base not in BASES_PICKUP or r.get("orderSourcename") != ORIGEM_PEDIDO_FILTRO:
            continue
        agregado[base]["deveria"] += r.get("shouldTakingNum", 0) or 0
        agregado[base]["coletado_no_prazo"] += r.get("timelyTakingNum", 0) or 0
        agregado[base]["tentativas"] += r.get("timelyTryTakingNum", 0) or 0

    faltando = sorted(set(BASES_PICKUP) - set(agregado.keys()))
    if faltando:
        log.warning("Bases da lista sem dados na API em %s: %s", dia, ", ".join(faltando))

    epop_brutos = _buscar_registros_epop_brutos(dia)
    epop_agregado: dict[str, dict[str, int]] = defaultdict(lambda: {"total": 0, "com_imagem": 0})
    for r in epop_brutos:
        base = (r.get("pickNetworkName") or "").strip()
        if base not in BASES_PICKUP or r.get("orderSourceName") != ORIGEM_PEDIDO_FILTRO:
            continue
        epop_agregado[base]["total"] += 1
        if r.get("signaturePictureUrl"):
            epop_agregado[base]["com_imagem"] += 1

    registros = []
    for base in BASES_PICKUP:
        valores = agregado.get(base)
        epop_valores = epop_agregado.get(base, {"total": 0, "com_imagem": 0})
        epop_total = epop_valores["total"]
        epop_com_imagem = epop_valores["com_imagem"]
        taxa_poc_pct = round(epop_com_imagem / epop_total * 100, 2) if epop_total else None

        if valores is None:
            registros.append({
                "base": base,
                "qtd_a_coletar": 0,
                "qtd_coletada_no_prazo": 0,
                "soma_coletados_tentativas": 0,
                "taxa_real_pct": None,
                "taxa_com_tentativas_pct": None,
                "epop_total": epop_total,
                "epop_com_imagem": epop_com_imagem,
                "taxa_poc_pct": taxa_poc_pct,
            })
            continue

        deveria = valores["deveria"]
        coletado_no_prazo = valores["coletado_no_prazo"]
        tentativas = valores["tentativas"]
        taxa_pct = round(coletado_no_prazo / deveria * 100, 2) if deveria else None
        taxa_com_tentativas_pct = round(tentativas / deveria * 100, 2) if deveria else None
        registros.append({
            "base": base,
            "qtd_a_coletar": deveria,
            "qtd_coletada_no_prazo": coletado_no_prazo,
            "soma_coletados_tentativas": tentativas,
            "taxa_real_pct": taxa_pct,
            "taxa_com_tentativas_pct": taxa_com_tentativas_pct,
            "epop_total": epop_total,
            "epop_com_imagem": epop_com_imagem,
            "taxa_poc_pct": taxa_poc_pct,
        })
    return registros


# Cabeçalhos bilíngues PT/CN (Guilherme, 2026-08-11) — mesmo estilo das
# planilhas que o JMS já exporta (ex.: "Base 基地", "TAXA 取件率"). O chinês
# das 6 primeiras colunas replica o que já apareceu nas planilhas de
# referência (print PICKUP e DROPOFF.xlsx); "Taxa de coleta com tentativas"
# e "Transferência POC" não tinham rótulo chinês capturado do portal — usei
# a tradução mais próxima do termo técnico (揽收及时率 é o nome do próprio
# indicador no JMS, visto na captura de rede do routerNameList).
CABECALHO = [
    "Data 日期",
    "Base 基地",
    "Qtd a coletar 应揽收数量",
    "Qtd coletada no prazo 及时揽收数量",
    "Soma de Pedidos coletados + Tentativas de coleta 已收订单总数 + 收款尝试次数",
    "Taxa de coleta Real 实际揽收准点率",
    "Taxa de coleta com tentativas de coleta 揽收及时率",
    "Transferência POC 签收图片",
]

COR_HEADER = "C00000"
COR_TENTATIVAS = "FFFF00"
PREENCHIMENTO_HEADER = PatternFill(start_color=COR_HEADER, end_color=COR_HEADER, fill_type="solid")
PREENCHIMENTO_TENTATIVAS = PatternFill(start_color=COR_TENTATIVAS, end_color=COR_TENTATIVAS, fill_type="solid")
FONTE_HEADER = Font(bold=True, color="FFFFFF")


def _formatar_aba(ws, num_linhas: int, num_colunas: int, cols_percentuais: list[int], col_tentativas: int):
    for cel in ws[1]:
        cel.font = FONTE_HEADER
        cel.fill = PREENCHIMENTO_HEADER
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    for col in cols_percentuais:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = "0.00%"
    for row in ws.iter_rows(min_row=2, min_col=col_tentativas, max_col=col_tentativas, max_row=num_linhas + 1):
        row[0].fill = PREENCHIMENTO_TENTATIVAS
    ultima_coluna = get_column_letter(num_colunas)
    ws.auto_filter.ref = f"A1:{ultima_coluna}{num_linhas + 1}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _linha_registro(r: dict) -> list:
    return [
        r["qtd_a_coletar"],
        r["qtd_coletada_no_prazo"],
        r["soma_coletados_tentativas"],
        (r["taxa_real_pct"] / 100) if r["taxa_real_pct"] is not None else None,
        (r["taxa_com_tentativas_pct"] / 100) if r["taxa_com_tentativas_pct"] is not None else None,
        (r["taxa_poc_pct"] / 100) if r["taxa_poc_pct"] is not None else None,
    ]


def _escrever_aba(ws, dia: str, registros: list[dict]):
    ws.append(CABECALHO)
    for r in registros:
        ws.append([dia, r["base"], *_linha_registro(r)])

    _formatar_aba(ws, len(registros), len(CABECALHO), cols_percentuais=[6, 7, 8], col_tentativas=7)
    larguras = [12, 14, 14, 20, 44, 18, 26, 18]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def montar_planilha(dia: str, registros: list[dict]) -> BytesIO:
    """Arquivo diário (a partir de amanhã): 1 aba só, com o dia informado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Taxa de Coleta no Prazo"
    _escrever_aba(ws, dia, registros)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def montar_planilha_multiplos_dias(dados: list[tuple[str, list[dict]]]) -> BytesIO:
    """Backfill: 1 aba por dia, na ordem em que os dias forem passados."""
    wb = Workbook()
    wb.remove(wb.active)
    for dia, registros in dados:
        ws = wb.create_sheet(title=dia)
        _escrever_aba(ws, dia, registros)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


CABECALHO_CONSOLIDADO = [
    "Base 基地",
    "Qtd a coletar 应揽收数量",
    "Qtd coletada no prazo 及时揽收数量",
    "Soma de Pedidos coletados + Tentativas de coleta 已收订单总数 + 收款尝试次数",
    "Taxa de coleta Real 实际揽收准点率",
    "Taxa de coleta com tentativas de coleta 揽收及时率",
    "Transferência POC 签收图片",
]


def _consolidar_por_base(dados: list[tuple[str, list[dict]]]) -> list[dict]:
    """Soma os números de cada base ao longo de todos os dias em `dados` e
    recalcula as taxas a partir dos totais (não é média das taxas diárias —
    é peso pelo volume, senão um dia de base pequena pesaria igual a um dia
    de base grande)."""
    somas: dict[str, dict[str, int]] = defaultdict(
        lambda: {"deveria": 0, "coletado_no_prazo": 0, "tentativas": 0, "epop_total": 0, "epop_com_imagem": 0}
    )
    for _dia, registros in dados:
        for r in registros:
            base = r["base"]
            somas[base]["deveria"] += r["qtd_a_coletar"]
            somas[base]["coletado_no_prazo"] += r["qtd_coletada_no_prazo"]
            somas[base]["tentativas"] += r["soma_coletados_tentativas"]
            somas[base]["epop_total"] += r["epop_total"]
            somas[base]["epop_com_imagem"] += r["epop_com_imagem"]

    vazio = {"deveria": 0, "coletado_no_prazo": 0, "tentativas": 0, "epop_total": 0, "epop_com_imagem": 0}
    consolidado = []
    for base in BASES_PICKUP:
        v = somas.get(base, vazio)
        taxa_pct = round(v["coletado_no_prazo"] / v["deveria"] * 100, 2) if v["deveria"] else None
        taxa_com_tentativas_pct = round(v["tentativas"] / v["deveria"] * 100, 2) if v["deveria"] else None
        taxa_poc_pct = round(v["epop_com_imagem"] / v["epop_total"] * 100, 2) if v["epop_total"] else None
        consolidado.append({
            "base": base,
            "qtd_a_coletar": v["deveria"],
            "qtd_coletada_no_prazo": v["coletado_no_prazo"],
            "soma_coletados_tentativas": v["tentativas"],
            "taxa_real_pct": taxa_pct,
            "taxa_com_tentativas_pct": taxa_com_tentativas_pct,
            "epop_total": v["epop_total"],
            "epop_com_imagem": v["epop_com_imagem"],
            "taxa_poc_pct": taxa_poc_pct,
        })
    return consolidado


def montar_painel_mensal(dados: list[tuple[str, list[dict]]]) -> BytesIO:
    """Painel mensal: aba 'Consolidado' (totais do período, 1 linha por
    base) + aba 'Geral Diário' (1 linha por dia x base, com autofiltro)."""
    wb = Workbook()

    ws_consolidado = wb.active
    ws_consolidado.title = "Consolidado"
    ws_consolidado.append(CABECALHO_CONSOLIDADO)
    consolidado = _consolidar_por_base(dados)
    for r in consolidado:
        ws_consolidado.append([r["base"], *_linha_registro(r)])
    _formatar_aba(ws_consolidado, len(consolidado), len(CABECALHO_CONSOLIDADO), cols_percentuais=[5, 6, 7], col_tentativas=6)
    for i, largura in enumerate([14, 14, 20, 44, 18, 26, 18], start=1):
        ws_consolidado.column_dimensions[get_column_letter(i)].width = largura

    ws_diario = wb.create_sheet(title="Geral Diário")
    ws_diario.append(CABECALHO)
    total_linhas = 0
    for dia, registros in dados:
        for r in registros:
            ws_diario.append([dia, r["base"], *_linha_registro(r)])
            total_linhas += 1
    _formatar_aba(ws_diario, total_linhas, len(CABECALHO), cols_percentuais=[6, 7, 8], col_tentativas=7)
    for i, largura in enumerate([12, 14, 14, 20, 44, 18, 26, 18], start=1):
        ws_diario.column_dimensions[get_column_letter(i)].width = largura

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def salvar_copia_local(nome_arquivo: str, conteudo: bytes) -> Path:
    caminho = PASTA_OUTPUT / nome_arquivo
    caminho.write_bytes(conteudo)
    return caminho


def enviar_email(dia: str, anexos: list[tuple[str, bytes]]):
    """`anexos`: lista de (nome_do_arquivo, conteudo_bytes) — manda todos
    juntos no mesmo e-mail (pedido do Guilherme, 2026-08-11: diário +
    painel mensal consolidado no mesmo envio)."""
    remetente = os.environ["EMAIL_SENDER"]
    senha_app = os.environ["EMAIL_APP_PASSWORD"]
    destinatarios = [e.strip() for e in os.environ["EMAIL_TO_PICKUP"].split(",") if e.strip()]

    msg = EmailMessage()
    msg["Subject"] = f"Taxa de coleta no prazo (PICKUP) - {dia}"
    msg["From"] = remetente
    msg["To"] = ", ".join(destinatarios)
    msg.set_content(
        f"Segue em anexo a taxa de coleta no prazo do dia {dia} "
        "(por base, origem TikTok, horário de término do prazo de coleta) "
        "e o painel consolidado do mês até essa data."
    )
    for nome_arquivo, conteudo in anexos:
        msg.add_attachment(
            conteudo,
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=nome_arquivo,
        )

    with smtplib.SMTP_SSL("smtp.feishu.cn", 465, context=ssl.create_default_context()) as smtp:
        smtp.login(remetente, senha_app)
        smtp.send_message(msg)


def main():
    dia = sys.argv[1] if len(sys.argv) > 1 else (date.today() - timedelta(days=1)).isoformat()
    dia_dt = date.fromisoformat(dia)
    inicio_mes = dia_dt.replace(day=1)

    log.info("Buscando Taxa de coleta no prazo (PICKUP) para %s (só esse dia)...", dia)
    registros_dia = buscar_pickup(dia)

    historico = carregar_historico()
    historico[dia] = registros_dia
    historico = {d: r for d, r in historico.items() if d >= inicio_mes.isoformat()}
    salvar_historico(historico)

    dados_mes = [(d, historico[d]) for d in sorted(historico.keys())]

    conteudo_diario = montar_planilha(dia, registros_dia).getvalue()
    conteudo_painel = montar_painel_mensal(dados_mes).getvalue()

    # Nome fixo (sem data) na cópia local/git — sobrescreve todo dia em vez
    # de acumular um arquivo por dia. O histórico já fica preservado nos
    # e-mails enviados (pedido do Guilherme, 2026-08-11: "não pesar" o
    # repositório com um arquivo novo por execução).
    caminho_diario = salvar_copia_local("PICKUP_diario.xlsx", conteudo_diario)
    caminho_painel = salvar_copia_local("PICKUP_painel_mensal.xlsx", conteudo_painel)
    log.info("Planilhas salvas em %s e %s", caminho_diario, caminho_painel)

    # No e-mail (o "histórico"), o nome do anexo leva a data — só a cópia
    # local/git é que fica com nome fixo.
    nome_anexo_diario = f"PICKUP_{dia}.xlsx"
    nome_anexo_painel = f"PICKUP_Painel_{inicio_mes.isoformat()[:7]}.xlsx"
    enviar_email(dia, [(nome_anexo_diario, conteudo_diario), (nome_anexo_painel, conteudo_painel)])
    log.info("E-mail enviado para %s", os.environ["EMAIL_TO_PICKUP"])


if __name__ == "__main__":
    try:
        main()
    except TokenExpiradoError as e:
        log.error(str(e))
        sys.exit(1)
